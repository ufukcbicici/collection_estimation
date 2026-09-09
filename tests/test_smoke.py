"""Smoke tests: the environment, the corpus copy and the config paths line up.

These are deliberately shallow. They exist so that a broken setup fails here, with a clear
message, rather than three modules deep into a modelling run.
"""
import json
import os
import sys

import pytest

sys.path.insert(0, os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))

from collection_estimation import config  # noqa: E402


def test_required_libraries_import():
    import numpy, pandas, sklearn, openpyxl, joblib  # noqa: F401


def test_corpus_directory_is_present_and_the_expected_size():
    assert os.path.isdir(config.CORPUS_DIR), (
        f"corpus missing at {config.CORPUS_DIR} — copy it from the cockpit repo")
    books = [f for f in os.listdir(config.CORPUS_DIR)
             if f.upper().endswith(".XLSX") and not f.startswith("~$")]
    # 306 files across 312 business days: a few days deliberately have no file at all,
    # which is NOT the same as a day with zero collections.
    assert len(books) == 306, f"expected 306 workbooks, found {len(books)}"


def test_oracle_parses_and_describes_this_corpus():
    with open(config.ORACLE, encoding="utf-8") as fh:
        oracle = json.load(fh)
    assert oracle["synthetic"] is True
    assert oracle["corpus_dir"] == os.path.basename(config.CORPUS_DIR)
    assert oracle["range"]["rows_total"] > 25_000


def test_a_workbook_opens_and_has_the_expected_shape():
    from openpyxl import load_workbook

    books = sorted(f for f in os.listdir(config.CORPUS_DIR)
                   if f.upper().endswith(".XLSX") and not f.startswith("~$"))
    wb = load_workbook(os.path.join(config.CORPUS_DIR, books[0]),
                       read_only=True, data_only=True)
    try:
        assert wb.sheetnames[0] == "_SYNTHETIC", (
            "the synthetic marker tab must be first — this data must never be "
            "mistaken for an EY record")
        ws = wb[wb.sheetnames[1]]
        grid = list(ws.iter_rows(min_row=1, max_row=config.HEADER_ROW,
                                 values_only=True))
        assert grid[0][0] == "MERCURY COLLECTIONS"
        header = grid[config.HEADER_ROW - 1]
        assert header[0] == "Company Code"
        assert header[3] == "Document Date"
        assert header[9] == "Amount in local currency"
    finally:
        wb.close()


@pytest.mark.parametrize("code", sorted(config.ENTITY_LABEL))
def test_every_entity_has_a_collection_group(code):
    assert code in config.ENTITY_GROUP
