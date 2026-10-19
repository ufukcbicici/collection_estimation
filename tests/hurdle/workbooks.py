"""Builders for small daily workbooks, shared by the hurdle reader tests.

`make_workbook` writes a workbook in the layout of the daily files: an optional marker sheet
`_SYNTHETIC`, one data sheet, a title on row 1, the header on row 2, payment lines from row 3,
and a grand-total row last. Customers are fake.
"""
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from collection_estimation.hurdle.holidays import HolidayTable
from collection_estimation.hurdle.settings import CollectionsLayout, LedgerLayout

HEADER = ('Company Code', 'Customer', 'Customer Name', 'Document Date', 'Bank',
          'SAP Document No', 'Government Inv', 'Amount in Original Currency',
          'Document currency', 'Amount in local currency', 'Local Currency')

# The layout of the workbooks `make_workbook` writes, as [collections_layout] in
# settings/synthetic.toml gives it: "02062025 ..." file names, the header on row 2, payment
# lines from row 3, and the 11 headers of HEADER.
SYNTHETIC_HEADERS = {
    'Company Code': 'company_code',
    'Customer': 'customer',
    'Customer Name': 'customer_name',
    'Document Date': 'document_date',
    'Bank': 'bank',
    'SAP Document No': 'sap_document_no',
    'Government Inv': 'invoice_no',
    'Amount in Original Currency': 'amount_original',
    'Document currency': 'currency_original',
    'Amount in local currency': 'amount_try',
    'Local Currency': 'currency_local',
}
SYNTHETIC_LAYOUT = CollectionsLayout(file_name_date_format='%d%m%Y', header_row=2,
                                     first_data_row=3, headers=SYNTHETIC_HEADERS)

# Column positions in HEADER, for building subtotal and grand-total rows.
NAME_COLUMN = HEADER.index('Customer Name')
AMOUNT_TRY_COLUMN = HEADER.index('Amount in local currency')
LOCAL_CURRENCY_COLUMN = HEADER.index('Local Currency')


def line(customer='0012345', amount_try=100.0, *, company='TR02', name='ACME TEST A.S.',
         when=datetime(2025, 6, 2), bank='TEST BANK', sap='1400000001', invoice='GIB2025000001',
         amount_original=None, currency='TRY'):
    """One payment line as a tuple in HEADER order."""
    if amount_original is None:
        amount_original = amount_try
    return (company, customer, name, when, bank, sap, invoice, amount_original, currency,
            amount_try, 'TRY')


def total_row(label, amount, width=len(HEADER)):
    """A subtotal or grand-total row: no company code or customer, a label and a TRY amount."""
    row = [None] * width
    row[NAME_COLUMN] = label
    row[AMOUNT_TRY_COLUMN] = amount
    row[LOCAL_CURRENCY_COLUMN] = 'TRY'
    return tuple(row)


GRAND_TOTAL_DEFAULT = object()  # "use the sum of the lines"


def make_workbook(folder, lines, grand_total=GRAND_TOTAL_DEFAULT, *, marker=True,
                  extra_sheets=(), last_rows=None, header=HEADER, name='02062025',
                  title_rows=(('MERCURY COLLECTIONS',),)):
    """Write a daily workbook to `folder` and return its path.

    `lines` are rows written from row 3 (payment lines, subtotal rows, or None for a blank
    spacer row). A grand-total row follows, `grand_total` defaulting to the sum of the payment
    lines' TRY amounts; `last_rows` replaces that grand-total row with rows of your choice.
    `title_rows` are the rows above the header (one title row by default; `()` puts the
    header on row 1 and the lines from row 2).
    """
    workbook = Workbook()
    first = workbook.active
    if marker:
        first.title = '_SYNTHETIC'
        sheet = workbook.create_sheet(name)
    else:
        first.title = name
        sheet = first
    for extra in extra_sheets:
        workbook.create_sheet(extra)

    for title in title_rows:
        sheet.append(list(title))
    sheet.append(list(header))
    for row in lines:
        if row is None:
            sheet.append([])  # openpyxl skips empty appends' cells; the row number still moves
        else:
            sheet.append(list(row))

    if last_rows is None:
        if grand_total is GRAND_TOTAL_DEFAULT:
            grand_total = sum(row[AMOUNT_TRY_COLUMN] for row in lines
                              if row is not None and row[0] is not None)
        last_rows = [total_row('GRAND TOTAL', grand_total)]
    for row in last_rows:
        sheet.append(list(row))

    path = Path(folder) / f'{name} EY TAHSILAT.XLSX'
    workbook.save(path)
    return path


# --- AR ledger (FBL5N layout) -------------------------------------------------------------
# A ledger workbook: an optional marker sheet `_SYNTHETIC`, then a sheet 'FBL5N' with the
# header on row 1 and one ledger line per row from row 2. Customers are fake.

LEDGER_HEADER = ('Company Code', 'Customer', 'Customer Name', 'Document Number',
                 'Document Type', 'Reference', 'Invoice Reference', 'Document Date',
                 'Posting Date', 'Baseline Date', 'Terms of Payment', 'Net Due Date',
                 'Amount in Doc. Curr.', 'Document Currency', 'Amount in Local Currency',
                 'Clearing Date', 'Clearing Document')

# The layout of the ledgers `make_ledger_workbook` writes, as [ledger_layout] in
# settings/synthetic.toml gives it: the header on row 1, ledger lines from row 2, the 17
# headers of LEDGER_HEADER, and RV / DZ / DG as invoice / payment / credit memo.
# 'Clearing Document' is a required column (its cells may be empty on open lines); only
# Customer Name, Document Date, Baseline Date and Terms of Payment may be left out.
SYNTHETIC_LEDGER_HEADERS = {
    'Company Code': 'company_code',
    'Customer': 'customer',
    'Customer Name': 'customer_name',
    'Document Number': 'document_no',
    'Document Type': 'document_type',
    'Reference': 'reference',
    'Invoice Reference': 'invoice_reference',
    'Document Date': 'document_date',
    'Posting Date': 'posting_date',
    'Baseline Date': 'baseline_date',
    'Terms of Payment': 'payment_terms',
    'Net Due Date': 'net_due_date',
    'Amount in Doc. Curr.': 'amount_document',
    'Document Currency': 'document_currency',
    'Amount in Local Currency': 'amount_local',
    'Clearing Date': 'clearing_date',
    'Clearing Document': 'clearing_document',
}
SYNTHETIC_DOCUMENT_TYPES = {'invoice': ('RV',), 'payment': ('DZ',), 'credit_memo': ('DG',)}
SYNTHETIC_LEDGER_LAYOUT = LedgerLayout(header_row=1, first_data_row=2,
                                       headers=SYNTHETIC_LEDGER_HEADERS,
                                       document_types=SYNTHETIC_DOCUMENT_TYPES)


def ledger_row(*, company='TR02', customer='0012345', name='ACME TEST A.S.',
               document_no='9000000001', document_type='RV', reference=None,
               invoice_reference=None, document_date=datetime(2026, 3, 2),
               posting_date=datetime(2026, 3, 2), baseline_date=None, terms=None,
               net_due_date=None, amount=100.0, currency='TRY', amount_local=None,
               clearing_date=None, clearing_document=None):
    """One ledger line as a dict keyed by LEDGER_HEADER text (amount_local defaults to amount)."""
    if amount_local is None:
        amount_local = amount
    values = (company, customer, name, document_no, document_type, reference,
              invoice_reference, document_date, posting_date, baseline_date, terms,
              net_due_date, amount, currency, amount_local, clearing_date, clearing_document)
    return dict(zip(LEDGER_HEADER, values))


def invoice(document_no='9000000001', amount=1000.0, *, e_number='GIB2026000001',
            posted=datetime(2026, 3, 2), due=datetime(2026, 4, 1), terms='Z030', **kwargs):
    """An invoice (RV): positive amount, a net due date and payment terms."""
    return ledger_row(document_no=document_no, document_type='RV', reference=e_number,
                      document_date=posted, posting_date=posted, baseline_date=posted,
                      terms=terms, net_due_date=due, amount=amount, **kwargs)


def payment(document_no='1400000001', amount=1000.0, *, invoice_no='9000000001',
            posted=datetime(2026, 3, 20), **kwargs):
    """A payment (DZ) applied to an invoice: negative amount, no due date."""
    return ledger_row(document_no=document_no, document_type='DZ', invoice_reference=invoice_no,
                      document_date=posted, posting_date=posted, amount=-abs(amount), **kwargs)


def unapplied_payment(document_no='1400000002', amount=500.0, *,
                      posted=datetime(2026, 3, 21), **kwargs):
    """A payment (DZ) not applied to any invoice: negative amount, no invoice reference."""
    return ledger_row(document_no=document_no, document_type='DZ', invoice_reference=None,
                      document_date=posted, posting_date=posted, amount=-abs(amount), **kwargs)


def credit_memo(document_no='9100000001', amount=200.0, *, e_number='GIB2026000001',
                posted=datetime(2026, 3, 10), **kwargs):
    """A credit memo (DG): negative amount, reference = the invoice's e-invoice number."""
    return ledger_row(document_no=document_no, document_type='DG', reference=e_number,
                      document_date=posted, posting_date=posted, amount=-abs(amount), **kwargs)


def refund(document_no='1400000003', amount=200.0, *, memo_no='9100000001',
           posted=datetime(2026, 3, 25), **kwargs):
    """A refund paid out to the customer (DZ): positive amount, invoice reference = the memo."""
    return ledger_row(document_no=document_no, document_type='DZ', invoice_reference=memo_no,
                      document_date=posted, posting_date=posted, amount=abs(amount), **kwargs)


def make_ledger_workbook(folder, lines, *, marker=True, header=LEDGER_HEADER, extra_sheets=(),
                         name='SYNTHETIC_ledger.xlsx', title_rows=()):
    """Write a ledger workbook to `folder` and return its path.

    `title_rows` are rows written above the header (none by default: the header on row 1).

    `lines` are dicts from the makers above (written in `header` order, so a shuffled or
    shortened header moves or drops the values with it; header entries not in a dict are left
    empty), tuples/lists written as they are, or None for an empty row.
    """
    workbook = Workbook()
    first = workbook.active
    if marker:
        first.title = '_SYNTHETIC'
        sheet = workbook.create_sheet('FBL5N')
    else:
        first.title = 'FBL5N'
        sheet = first
    for extra in extra_sheets:
        workbook.create_sheet(extra)

    for title in title_rows:
        sheet.append(list(title))
    sheet.append(list(header))
    for row in lines:
        if row is None:
            sheet.append([])
        elif isinstance(row, dict):
            sheet.append([row.get(text) for text in header])
        else:
            sheet.append(list(row))

    path = Path(folder) / name
    workbook.save(path)
    return path


# --- public holidays ----------------------------------------------------------------------
# The Turkish holidays of settings/holidays_tr.toml, written out literally so that the tests
# do not depend on that file (test_holidays.py checks that the two are equal).
SYNTHETIC_HOLIDAYS = HolidayTable(
    fixed={
        (1, 1): "New Year's Day",
        (4, 23): "National Sovereignty and Children's Day",
        (5, 1): "Labour and Solidarity Day",
        (5, 19): "Commemoration of Atatürk, Youth and Sports Day",
        (7, 15): "Democracy and National Unity Day",
        (8, 30): "Victory Day",
        (10, 29): "Republic Day",
    },
    bayrams=(
        (date(2025, 3, 30), date(2025, 4, 1), "Ramazan Bayramı"),
        (date(2025, 6, 6), date(2025, 6, 9), "Kurban Bayramı"),
        (date(2026, 3, 20), date(2026, 3, 22), "Ramazan Bayramı"),
        (date(2026, 5, 27), date(2026, 5, 30), "Kurban Bayramı"),
    ),
)
