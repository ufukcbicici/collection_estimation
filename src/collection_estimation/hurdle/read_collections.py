"""Reading the daily collection workbooks (one .xlsx per business day).

Gives the rest of the package two things: the payment lines, and the list of days that have
a file. A day without a file is MISSING, not a day without payments.
"""
import math
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .cache import file_fingerprint, load_with_cache, source_fingerprint

# "02062025 EY TAHSILAT.XLSX" is 2 June 2025. Change here if the real files are named otherwise.
FILE_NAME_DATE_FORMAT = "%d%m%Y"
FILE_NAME_DATE_LENGTH = 8

# Header text in the workbook -> our field name. Headers are compared after folding case
# and extra spaces (see `simplify_header`). If the real files use Turkish headers, add the
# Turkish text as further keys pointing to the same field names.
HEADER_TO_FIELD = {
    "company code": "company_code",
    "customer": "customer",
    "customer name": "customer_name",
    "document date": "document_date",
    "bank": "bank",
    "sap document no": "sap_document_no",
    "government inv": "invoice_no",
    "amount in original currency": "amount_original",
    "document currency": "currency_original",
    "amount in local currency": "amount_try",
    "local currency": "currency_local",
}
OPTIONAL_FIELDS = {"bank"}  # the real files may not carry a bank column

# Row 1 of a daily file is a title ("MERCURY COLLECTIONS"), row 2 the header.
HEADER_ROW = 2
FIRST_DATA_ROW = 3

# The payment lines' grand total must equal the file's own grand total within this (TRY).
TIE_OUT_TOLERANCE = 0.01

# The columns of the table `read_one_file` returns, in order.
LINE_COLUMNS = list(HEADER_TO_FIELD.values()) + ["source_file", "source_row"]

# The text columns of that table: everything except the date, the amounts and the row number.
TEXT_COLUMNS = [column for column in LINE_COLUMNS
                if column not in ("document_date", "amount_original", "amount_try", "source_row")]


def list_daily_files(folder: Path) -> dict[date, Path]:
    """Every daily workbook in `folder`, keyed by the date in its name, oldest first."""
    files = {}
    for path in Path(folder).iterdir():
        if not path.is_file():
            continue  # e.g. a sub-folder, even one named like a workbook
        if path.suffix.lower() != ".xlsx":
            continue
        if path.name.startswith("~$"):
            continue  # Excel's lock file for an open workbook, not a real file
        file_date = date_from_file_name(path.name)
        if file_date in files:
            raise ValueError(
                f"two files for {file_date}: {files[file_date].name} and {path.name}")
        files[file_date] = path
    # Sort by the date itself: sorting the names would put 01072026 before 02062025.
    return dict(sorted(files.items()))


def date_from_file_name(name: str) -> date:
    """The date in the first 8 characters of a file name, e.g. '02062025' -> 2 June 2025."""
    try:
        return datetime.strptime(name[:FILE_NAME_DATE_LENGTH], FILE_NAME_DATE_FORMAT).date()
    except ValueError:
        raise ValueError(
            f"{name}: the file name does not start with a date in the form DDMMYYYY") from None


def find_columns(header_row: tuple, header_to_field: dict[str, str] | None = None,
                 optional_fields: set[str] | None = None) -> dict[str, int]:
    """{field name: zero-based column position}, matched on the header text.

    Columns are found by their header text, never by their position: if a column is ever
    inserted into a file, reading by position would silently shift every column to its
    right. `header_to_field` and `optional_fields` describe the file; they default to the
    daily collection files' HEADER_TO_FIELD and OPTIONAL_FIELDS (the AR ledger passes its own).
    """
    if header_to_field is None:
        header_to_field = HEADER_TO_FIELD
    if optional_fields is None:
        optional_fields = OPTIONAL_FIELDS

    # Simplify the dictionary's keys the same way as the file's headers, so a key matches
    # however it was typed (e.g. 'MÜŞTERİ' or 'müşteri').
    field_by_header = {simplify_header(text): field for text, field in header_to_field.items()}

    columns = {}
    for position, header in enumerate(header_row):
        field = field_by_header.get(simplify_header(header))
        if field is None:
            continue  # a column we don't use
        if field in columns:
            first, second = get_column_letter(columns[field] + 1), get_column_letter(position + 1)
            raise ValueError(f"field '{field}' is in both column {first} and column {second}")
        columns[field] = position

    required = set(header_to_field.values()) - optional_fields
    missing = required - set(columns)
    if missing:
        raise ValueError(
            f"missing columns {sorted(missing)}; the header row reads {header_row}")
    return columns


def simplify_header(text) -> str:
    """'  Company  CODE ' -> 'company code'; 'MÜŞTERİ' -> 'müşteri'; an empty cell -> ''."""
    if text is None:
        return ""
    simple = " ".join(str(text).split()).casefold()
    # casefold() turns the Turkish 'İ' into 'i' followed by a separate combining dot
    # (U+0307); drop the dot so that 'MÜŞTERİ' and 'müşteri' match.
    return simple.replace("̇", "")


def read_one_file(path: Path) -> pd.DataFrame:
    """The payment lines of one daily workbook, one row per line, in sheet order.

    Nothing is dropped and nothing is computed: refunds (negative amounts) and unapplied cash
    stay in. Subtotal, grand-total and blank rows are left out; they have no company code.
    Raises ValueError if the file is malformed or its lines do not add up to its grand total.
    """
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[pick_data_sheet(workbook.sheetnames, path.name)]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if len(rows) < HEADER_ROW:
        raise ValueError(f"{path.name}: no header row (row {HEADER_ROW})")
    columns = find_columns(rows[HEADER_ROW - 1])

    lines = []
    for row_number, row in enumerate(rows[FIRST_DATA_ROW - 1:], start=FIRST_DATA_ROW):
        if is_blank(cell(row, columns, "company_code")):
            continue  # a subtotal, the grand total, or an empty spacer row
        lines.append(payment_line(row, columns, path.name, row_number))

    table = pd.DataFrame(lines, columns=LINE_COLUMNS)
    # Set the column types explicitly: a file with no payment lines would otherwise give
    # columns of generic objects, which then spread to the table of all days.
    table["document_date"] = pd.to_datetime(table["document_date"])
    table["amount_original"] = table["amount_original"].astype("float64")
    table["amount_try"] = table["amount_try"].astype("float64")
    table["source_row"] = table["source_row"].astype("int64")
    for column in TEXT_COLUMNS:
        # pandas stores a missing text value as NaN or None depending on the rest of the
        # column; make it always None, so later code can simply test `is None`.
        table[column] = table[column].astype(object).where(table[column].notna(), None)
    check_grand_total(table, rows, columns, path.name)
    return table


def pick_data_sheet(sheet_names: list[str], file_name: str) -> str:
    """The one sheet holding the data. Sheets whose name starts with '_' are markers."""
    data_sheets = [name for name in sheet_names if not name.startswith("_")]
    if len(data_sheets) != 1:
        raise ValueError(
            f"{file_name}: expected exactly one data sheet, found {data_sheets}")
    return data_sheets[0]


def cell(row: tuple, columns: dict[str, int], field: str):
    """The value of `field` in `row`; None if the file has no such column or the row is short."""
    position = columns.get(field)
    if position is None or position >= len(row):
        return None
    return row[position]


def is_blank(value) -> bool:
    """True for an empty cell or a cell holding only spaces."""
    return value is None or str(value).strip() == ""


def payment_line(row: tuple, columns: dict[str, int], file_name: str, row_number: int) -> dict:
    """One payment line as a dict, after checking the type of every field we rely on."""
    where = f"{file_name} row {row_number}"

    customer = cell(row, columns, "customer")
    if not isinstance(customer, str):
        # Read as a number, a customer ID has already lost its leading zeros.
        raise ValueError(f"{where}: customer is {customer!r}, not text")

    document_date = cell(row, columns, "document_date")
    if not isinstance(document_date, datetime):  # Excel dates arrive as datetime
        raise ValueError(f"{where}: document date is {document_date!r}, not a date")

    line = {field: cell(row, columns, field) for field in HEADER_TO_FIELD.values()}
    line["company_code"] = str(line["company_code"]).strip()
    line["customer"] = customer.strip()
    line["document_date"] = document_date.date()

    for field in ("amount_original", "amount_try"):
        amount = line[field]
        # bool is a subclass of int in Python, so it is excluded explicitly.
        if isinstance(amount, bool) or not isinstance(amount, (int, float)):
            raise ValueError(f"{where}: {field} is {amount!r}, not a number")
        # NaN counts as a float in Python, but later sums would silently skip it.
        if is_nan(amount):
            raise ValueError(f"{where}: {field} is NaN, not a number")
        line[field] = float(amount)

    line["source_file"] = file_name
    line["source_row"] = row_number
    return line


def check_grand_total(lines: pd.DataFrame, rows: list[tuple], columns: dict[str, int],
                      file_name: str) -> None:
    """Check that the payment lines add up to the file's grand total.

    The grand total is the last row that is not empty. This catches lines the reader skipped
    by mistake. The check is skipped with a note if the file has no grand-total row (its last
    row is a payment line) or the total has no stored number (an Excel formula never
    calculated). Any other content in that row raises ValueError.
    """
    total_row = last_non_empty_row(rows)
    grand_total = cell(total_row, columns, "amount_try")

    if not is_blank(cell(total_row, columns, "company_code")):
        print(f"Note: {file_name} has no grand-total row; tie-out skipped")
        return
    if grand_total is None or is_nan(grand_total):
        print(f"Note: {file_name} has no number in its grand total; tie-out skipped")
        return
    # bool is a subclass of int in Python, so it is excluded explicitly.
    if isinstance(grand_total, bool) or not isinstance(grand_total, (int, float)):
        raise ValueError(f"{file_name}: the grand total is {grand_total!r}, not a number")

    lines_total = lines["amount_try"].sum()
    if abs(lines_total - grand_total) > TIE_OUT_TOLERANCE:
        raise ValueError(
            f"{file_name}: the payment lines add up to {lines_total:,.2f} TRY "
            f"but the file's grand total is {grand_total:,.2f} TRY")


def last_non_empty_row(rows: list[tuple]) -> tuple:
    """The last row with at least one non-empty cell (stray empty rows at the end are skipped)."""
    for row in reversed(rows):
        if not all(is_blank(value) for value in row):
            return row
    return ()


def is_nan(value) -> bool:
    """True for a float NaN ('not a number'), which would make every comparison with it False."""
    return isinstance(value, float) and math.isnan(value)


def read_all_files(daily_files: dict[date, Path]) -> pd.DataFrame:
    """The payment lines of every daily file, oldest day first, in sheet order within a day.

    `daily_files` is the output of `list_daily_files`; its keys are the days that have a file.
    """
    if not daily_files:
        raise ValueError("no daily files to read")

    tables = []
    for count, (file_date, path) in enumerate(daily_files.items(), start=1):
        lines = read_one_file(path)
        check_dates_match(lines, file_date, path.name)
        tables.append(lines)
        if count % 50 == 0:
            print(f"  ... read {count} of {len(daily_files)} files")
    return pd.concat(tables, ignore_index=True)


def check_dates_match(lines: pd.DataFrame, file_date: date, file_name: str) -> None:
    """Every line's document date must be the date in its file's name.

    A file holding another day's lines would put that cash in the wrong week.
    """
    line_dates = set(lines["document_date"].dt.date)
    wrong_dates = sorted(line_dates - {file_date})
    if wrong_dates:
        listed = ", ".join(str(day) for day in wrong_dates)
        raise ValueError(
            f"{file_name}: the file name says {file_date} but lines carry {listed}")


def load_collections(daily_files: dict[date, Path], cache_folder: Path) -> pd.DataFrame:
    """The payment lines of every daily file, from the cache if nothing has changed.

    The cache is read again automatically if a file is added, removed or changed, if the
    files are in another folder, if this module's code changes, or if the cache is damaged.
    """
    fingerprint = [files_fingerprint(daily_files), source_fingerprint(Path(__file__))]
    return load_with_cache("hurdle_collections", fingerprint,
                           lambda: read_all_files(daily_files), cache_folder)


def files_fingerprint(daily_files: dict[date, Path]) -> list[list]:
    """[full path, size, last-modified time] for each daily file, in date order."""
    return [file_fingerprint(path) for path in daily_files.values()]
