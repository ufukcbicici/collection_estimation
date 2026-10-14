"""Reading the AR (accounts receivable) ledger: every invoice, payment and credit memo on each
customer's account, open and cleared, in the layout of SAP's customer line-item report (FBL5N).

The field names follow `ar_ledger_schema.md` in the cockpit repo. Working out what was open
at each past week is a later step; this module only reads and checks.

Sign convention (SAP's): what the customer owes is positive. An invoice is +, a payment -,
a credit memo -, and a refund paid out to the customer +.
"""
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from . import read_collections
from .cache import file_fingerprint, load_with_cache, source_fingerprint
from .read_collections import cell, find_columns, is_blank, is_nan, pick_data_sheet

# FBL5N header text -> our field name. Compared after folding case and spaces. When the real
# ledger arrives, add its labels as further keys pointing to the same field names.
LEDGER_HEADER_TO_FIELD = {
    "company code": "company_code",
    "customer": "customer",
    "customer name": "customer_name",
    "document number": "document_no",
    "document type": "document_type",
    "reference": "reference",
    "invoice reference": "invoice_reference",
    "document date": "document_date",
    "posting date": "posting_date",
    "baseline date": "baseline_date",
    "terms of payment": "payment_terms",
    "net due date": "net_due_date",
    "amount in doc. curr.": "amount_document",
    "document currency": "document_currency",
    "amount in local currency": "amount_local",
    "clearing date": "clearing_date",
    "clearing document": "clearing_document",
}
# Kept when present, for reconciliation; the model does not need them.
LEDGER_OPTIONAL_FIELDS = {
    "customer_name", "document_date", "baseline_date", "payment_terms", "clearing_document",
}

# The FBL5N sheet has its header on row 1.
LEDGER_HEADER_ROW = 1
LEDGER_FIRST_DATA_ROW = 2

# RV = invoice, DZ = payment (or a refund paid out), DG = credit memo.
DOCUMENT_TYPES = {"RV", "DZ", "DG"}

# The columns of the table `read_ledger` returns, in order, and which of them hold what.
LEDGER_COLUMNS = list(LEDGER_HEADER_TO_FIELD.values()) + ["source_row"]
LEDGER_DATE_COLUMNS = ["document_date", "posting_date", "baseline_date", "net_due_date",
                       "clearing_date"]
LEDGER_AMOUNT_COLUMNS = ["amount_document", "amount_local"]
LEDGER_TEXT_COLUMNS = [column for column in LEDGER_COLUMNS
                       if column not in LEDGER_DATE_COLUMNS + LEDGER_AMOUNT_COLUMNS
                       and column != "source_row"]
# Text fields every line must have.
REQUIRED_TEXT = ["company_code", "customer", "document_no", "document_type",
                 "document_currency"]


def load_ledger(path: Path, as_of: date, cache_folder: Path) -> pd.DataFrame:
    """The ledger lines (see `read_ledger`), from the cache if nothing has changed.

    The cache is read again automatically if the ledger file changes, `as_of` changes, the
    code of this module or of read_collections (whose helpers it uses) changes, or the cache
    is damaged.
    """
    fingerprint = [
        file_fingerprint(path),
        pd.Timestamp(as_of).date().isoformat(),  # '2026-08-28', whether a date or datetime
        source_fingerprint(Path(__file__), Path(read_collections.__file__)),
    ]
    return load_with_cache("hurdle_ledger", fingerprint,
                           lambda: read_ledger(path, as_of), cache_folder)


def read_ledger(path: Path, as_of: date) -> pd.DataFrame:
    """The ledger lines, one row per line, in sheet order, with our field names.

    `as_of` is the date the ledger was extracted; nothing in it may be dated later.
    Raises ValueError if the file is malformed, a line has a wrong type, a document type is
    unknown, an invoice has no due date, or a date is after `as_of`.
    """
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[pick_data_sheet(workbook.sheetnames, path.name)]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if len(rows) < LEDGER_HEADER_ROW:
        raise ValueError(f"{path.name}: no header row (row {LEDGER_HEADER_ROW})")
    columns = find_columns(rows[LEDGER_HEADER_ROW - 1], LEDGER_HEADER_TO_FIELD,
                           LEDGER_OPTIONAL_FIELDS)

    lines = []
    for row_number, row in enumerate(rows[LEDGER_FIRST_DATA_ROW - 1:],
                                     start=LEDGER_FIRST_DATA_ROW):
        if all(is_blank(value) for value in row):
            continue  # an empty row, e.g. at the end of the sheet
        lines.append(ledger_line(row, columns, f"{path.name} row {row_number}", row_number))

    ledger = pd.DataFrame(lines, columns=LEDGER_COLUMNS)
    # Set the column types explicitly, as for the daily files: an empty table would
    # otherwise give columns of generic objects.
    for column in LEDGER_DATE_COLUMNS:
        ledger[column] = pd.to_datetime(ledger[column])
    for column in LEDGER_AMOUNT_COLUMNS:
        ledger[column] = ledger[column].astype("float64")
    ledger["source_row"] = ledger["source_row"].astype("int64")
    for column in LEDGER_TEXT_COLUMNS:
        ledger[column] = ledger[column].astype(object).where(ledger[column].notna(), None)

    check_document_types(ledger, path.name)
    check_invoices_have_due_dates(ledger, path.name)
    check_nothing_after(ledger, as_of, path.name)
    return ledger


def ledger_line(row: tuple, columns: dict[str, int], where: str, row_number: int) -> dict:
    """One ledger line as a dict, after checking the type of every field.

    Text fields: text or empty (empty -> None). Dates: a date or empty, but the posting
    date is required. Amounts: a number, never NaN.
    """
    line = {}
    for field in LEDGER_HEADER_TO_FIELD.values():
        value = cell(row, columns, field)
        if field in LEDGER_DATE_COLUMNS:
            line[field] = date_or_none(value, where, field)
        elif field in LEDGER_AMOUNT_COLUMNS:
            line[field] = number(value, where, field)
        else:
            line[field] = text_or_none(value, where, field)

    for field in REQUIRED_TEXT:
        if line[field] is None:
            raise ValueError(f"{where}: {field} is empty")
    if line["posting_date"] is None:
        raise ValueError(f"{where}: posting_date is empty")
    line["document_type"] = line["document_type"].upper()
    line["source_row"] = row_number
    return line


def text_or_none(value, where: str, field: str):
    """A text cell with surrounding spaces removed; None for an empty cell.

    A number raises: read as a number, an ID such as a customer or document number may
    already have lost its leading zeros.
    """
    if is_blank(value):
        return None
    if not isinstance(value, str):
        raise ValueError(f"{where}: {field} is {value!r}, not text")
    return value.strip()


def date_or_none(value, where: str, field: str):
    """A date cell as a date (Excel dates arrive as datetime); None for an empty cell."""
    if is_blank(value):
        return None
    if not isinstance(value, datetime):
        raise ValueError(f"{where}: {field} is {value!r}, not a date")
    return value.date()


def number(value, where: str, field: str) -> float:
    """A number cell as a float. Empty, text, True/False and NaN raise."""
    # bool is a subclass of int in Python, so it is excluded explicitly.
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{where}: {field} is {value!r}, not a number")
    if is_nan(value):
        raise ValueError(f"{where}: {field} is NaN, not a number")
    return float(value)


def check_document_types(ledger: pd.DataFrame, file_name: str) -> None:
    """Raise ValueError naming every document type other than RV, DZ and DG, with its count.

    A real export may hold other types (e.g. opening balances, reversals); each one needs a
    decision before it is used, rather than being handled wrongly without notice.
    """
    is_unknown = ~ledger["document_type"].isin(DOCUMENT_TYPES)
    if is_unknown.any():
        lines_per_type = ledger.loc[is_unknown, "document_type"].value_counts().to_dict()
        raise ValueError(f"{file_name}: unknown document types (type: number of lines): "
                         f"{lines_per_type}; known types are {sorted(DOCUMENT_TYPES)}")


def check_invoices_have_due_dates(ledger: pd.DataFrame, file_name: str) -> None:
    """Raise ValueError if an invoice (RV) has no net due date."""
    no_due_date = (ledger["document_type"] == "RV") & ledger["net_due_date"].isna()
    if no_due_date.any():
        first_rows = ledger.loc[no_due_date, "source_row"].head(5).tolist()
        raise ValueError(f"{file_name}: {no_due_date.sum()} invoices have no net due date "
                         f"(first rows: {first_rows})")


def check_nothing_after(ledger: pd.DataFrame, as_of: date, file_name: str) -> None:
    """Raise ValueError if a posting or clearing date is after `as_of`, the extract date."""
    as_of_timestamp = pd.Timestamp(as_of)
    for column in ("posting_date", "clearing_date"):
        is_later = ledger[column] > as_of_timestamp
        if is_later.any():
            first_rows = ledger.loc[is_later, "source_row"].head(5).tolist()
            raise ValueError(f"{file_name}: {is_later.sum()} lines have a {column} after "
                             f"the extract date {as_of} (first rows: {first_rows})")


def check_ledger_against_collections(ledger: pd.DataFrame, collections: pd.DataFrame,
                                     days_with_file: list[date]) -> dict:
    """Check that the ledger's payments and the daily files describe the same money.

    For every day that has a daily file, the ledger's payment lines (DZ, sign flipped so
    money received is positive) posted that day must add up to the daily file's TRY total,
    within 0.01 TRY. Raises ValueError listing the first days that differ.

    Ledger payments on days without a file are expected and are not compared. They are
    counted in three groups: days before the first daily file and days after the last one
    (outside what the daily files cover, e.g. a ledger extracted later), and days in between
    (a missing file: cash the daily files lost). Returns those counts and the TRY they hold.
    """
    payments = ledger[ledger["document_type"] == "DZ"]
    received = -payments["amount_local"]  # the ledger records a payment received as negative
    ledger_per_day = received.groupby(payments["posting_date"].dt.date).sum()
    files_per_day = collections.groupby(collections["document_date"].dt.date)["amount_try"].sum()

    days = sorted({pd.Timestamp(day).date() for day in days_with_file})
    if not days:
        raise ValueError("no days with a file to compare")
    comparison = pd.DataFrame({
        "daily_files": files_per_day.reindex(days, fill_value=0.0),
        "ledger": ledger_per_day.reindex(days, fill_value=0.0),
    })
    differs = (comparison["daily_files"] - comparison["ledger"]).abs() > 0.01
    if differs.any():
        first_days = [f"{day}: daily files {row.daily_files:,.2f} TRY, ledger {row.ledger:,.2f} TRY"
                      for day, row in comparison[differs].head(5).iterrows()]
        raise ValueError(f"the ledger's payments differ from the daily files on {differs.sum()} "
                         f"days; first: {first_days}")

    without_file = ledger_per_day[~ledger_per_day.index.isin(days)]
    before_first_file = without_file[without_file.index < days[0]]
    after_last_file = without_file[without_file.index > days[-1]]
    inside_range = without_file[(without_file.index > days[0]) & (without_file.index < days[-1])]
    return {
        "days_compared": len(days),
        "days_before_first_file": len(before_first_file),
        "try_before_first_file": float(before_first_file.sum()),
        "missing_file_days_in_ledger": len(inside_range),
        "try_on_missing_file_days": float(inside_range.sum()),
        "days_after_last_file": len(after_last_file),
        "try_after_last_file": float(after_last_file.sum()),
    }
