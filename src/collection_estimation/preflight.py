"""Diagnose a workbook against what `config.py` declares — before running the pipeline.

    python src/collection_estimation/preflight.py  <folder-or-file>

`ingest.py` is deliberately strict: a file that does not match the expected layout raises
`WorkbookFormatError` rather than being quietly skipped, because a parser that silently
drops files fails as an unexplained dip in a weekly total months later.

That strictness is right for a run and unhelpful for a first look, where it tells you only
the *first* thing that disagrees. This module reports **everything** that disagrees, in one
pass, and never raises. Its output is a to-do list for `config.py`.

Nothing here is used by the pipeline. It is a diagnostic tool for a human.
"""
from __future__ import annotations

import os
import sys
from collections import Counter
from datetime import date, datetime

if __name__ == "__main__" and __package__ is None:      # run as a plain script
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    __package__ = "collection_estimation"

from openpyxl import load_workbook                                      # noqa: E402

from . import config                                                    # noqa: E402
from .ingest import REQUIRED_FIELDS, _normalise                         # noqa: E402

OK, WARN, BAD = "  ok  ", " warn ", " FAIL "


def _line(status: str, text: str) -> None:
    print(f"[{status}] {text}")


def _note(text: str) -> None:
    """A continuation line under a finding — advice, not a verdict."""
    print(f"         {text}")


# --------------------------------------------------------------------------------------
# The folder
# --------------------------------------------------------------------------------------

def inspect_folder(directory: str) -> list[str]:
    """Report on the set of files, and return the workbook paths in date order."""
    print(f"\n=== folder: {directory}")
    if not os.path.isdir(directory):
        _line(BAD, "not a directory")
        return []

    names = sorted(os.listdir(directory))
    workbooks = [n for n in names if n.upper().endswith(".XLSX")
                 and not n.startswith("~$")]
    locks = [n for n in names if n.startswith("~$")]
    others = [n for n in names
              if not n.upper().endswith(".XLSX") and not n.startswith("~$")]

    _line(OK if workbooks else BAD, f"{len(workbooks)} .xlsx files")
    if locks:
        _line(WARN, f"{len(locks)} Excel lock file(s) (~$...) — ignored, but it means a "
                    "workbook is open somewhere")
    if others:
        _line(WARN, f"{len(others)} non-xlsx entries, ignored: {others[:5]}")
    if not workbooks:
        return []

    # Filenames must encode their date, and are sorted on it rather than lexically.
    dated: list[tuple[date, str]] = []
    unparsed: list[str] = []
    for name in workbooks:
        head = name[:config.FILENAME_DATE_LENGTH]
        try:
            dated.append((datetime.strptime(
                head, config.FILENAME_DATE_FORMAT).date(), name))
        except ValueError:
            unparsed.append(name)

    if unparsed:
        _line(BAD, f"{len(unparsed)} filename(s) do not start with a date in "
                   f"'{config.FILENAME_DATE_FORMAT}': {unparsed[:5]}")
        _note("-> set config.FILENAME_DATE_FORMAT / FILENAME_DATE_LENGTH")
    else:
        _line(OK, f"all filenames parse as {config.FILENAME_DATE_FORMAT}")

    if dated:
        dated.sort()
        first, last = dated[0][0], dated[-1][0]
        span_weeks = ((last - first).days // 7) + 1
        _line(OK, f"date range {first} .. {last}  (~{span_weeks} weeks, "
                  f"{len(dated)} files)")
        if span_weeks < 30:
            _line(WARN, f"only ~{span_weeks} weeks — the rolling-origin evaluation "
                        "warms up on 12 and needs a good number of origins after that")

        # Business days with no file. On real data this is the honest version of the
        # synthetic corpus's oracle: a missing day is an ABSENT observation, not a zero.
        have = {d for d, _ in dated}
        expected = []
        day = first
        while day <= last:
            if day.weekday() < 5:
                expected.append(day)
            day = date.fromordinal(day.toordinal() + 1)

        # The holiday table has to be checked BEFORE the gaps can be interpreted: an
        # undeclared year makes every Bayram look like a missing file.
        years = sorted({d.year for d, _ in dated})
        from .calendar_features import covered_years, holiday_dates
        uncovered = [y for y in years if y not in covered_years()]
        if uncovered:
            _line(BAD, f"data covers {years}, but the moving-holiday table declares "
                       f"only {sorted(covered_years())}")
            _note(f"-> extend calendar_features.MOVING_HOLIDAYS for {uncovered}, or the")
            _note("   pipeline raises HolidayCoverageError. It raises ON PURPOSE: an")
            _note("   undeclared Bayram reads as five ordinary business days and")
            _note("   nothing complains.")
        else:
            _line(OK, f"moving-holiday table covers every year in the data {years}")

        # Weekdays with no file, split into the two cases that mean different things.
        holidays = holiday_dates(first, last, strict=False)
        gaps = [d for d in expected if d not in have]
        explained = [d for d in gaps if d in holidays]
        unexplained = [d for d in gaps if d not in holidays]

        if explained:
            _line(OK, f"{len(explained)} weekday(s) with no file are declared public "
                      f"holidays — expected")
        if unexplained:
            _line(WARN, f"{len(unexplained)} weekday(s) have no file and are NOT a "
                        f"declared holiday:")
            _note(f"   {[str(d) for d in unexplained[:8]]}")
            _note("-> either a holiday missing from the table, or a genuinely absent")
            _note("   file. They are DIFFERENT: a missing file is an absent")
            _note("   observation, not zero cash. The pipeline keeps the weeks that")
            _note("   contain one rather than dropping them.")
        if not gaps:
            _line(OK, "every weekday in range has a file")

    return [os.path.join(directory, n) for _, n in dated]


# --------------------------------------------------------------------------------------
# One workbook
# --------------------------------------------------------------------------------------

def inspect_workbook(path: str) -> None:
    """Report one workbook's layout against the constants in `config.py`."""
    print(f"\n=== workbook: {os.path.basename(path)}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:                              # noqa: BLE001 — report, not raise
        _line(BAD, f"cannot open: {type(exc).__name__}: {exc}")
        return

    try:
        _line(OK, f"sheets: {workbook.sheetnames}")

        sheet_name = None
        for name in workbook.sheetnames:
            stripped = name.strip()
            if (len(stripped) == config.FILENAME_DATE_LENGTH
                    and stripped.isdigit()):
                sheet_name = name
                break
        if sheet_name is None:
            for name in workbook.sheetnames:
                if _normalise(name) not in config.MARKER_TABS:
                    sheet_name = name
                    break
        if sheet_name is None:
            _line(BAD, "no usable data sheet — every tab is a known marker tab")
            return
        _line(OK, f"data sheet chosen: {sheet_name!r}")

        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(min_row=1, max_row=200, values_only=True))
        if not rows:
            _line(BAD, "sheet is empty")
            return

        _inspect_header(rows)
        _inspect_data(rows)
    finally:
        workbook.close()


def _inspect_header(rows: list[tuple]) -> None:
    """The header row: what matched, what did not, and what is missing."""
    if len(rows) < config.HEADER_ROW:
        _line(BAD, f"fewer than {config.HEADER_ROW} rows — no header at row "
                   f"{config.HEADER_ROW}")
        return

    header = rows[config.HEADER_ROW - 1]
    texts = [(i, _normalise(c)) for i, c in enumerate(header) if _normalise(c)]
    _line(OK, f"header row {config.HEADER_ROW} holds {len(texts)} non-empty cells")

    matched, unmatched = {}, []
    for index, text in texts:
        field = config.FIELD_BY_HEADER.get(text)
        if field is not None and field not in matched:
            matched[field] = index
        else:
            unmatched.append(text)

    missing = sorted(REQUIRED_FIELDS - set(matched))
    if missing:
        _line(BAD, f"missing required column(s): {missing}")
        _note(f"   header actually reads: {[t for _, t in texts]}")
        _note("-> add the real spelling as a key in config.FIELD_BY_HEADER, mapping")
        _note("   to the same field name. Keep the English key too.")
    else:
        _line(OK, f"all {len(REQUIRED_FIELDS)} required columns found")

    for field in sorted(config.OPTIONAL_FIELDS):
        if field not in matched:
            _line(WARN, f"optional column '{field}' absent — tolerated")

    if unmatched:
        _line(WARN, f"{len(unmatched)} header cell(s) not recognised (ignored): "
                    f"{unmatched[:8]}")

    # A header on the wrong row is a common and confusing failure, so look for one.
    if missing:
        for candidate in range(1, min(6, len(rows) + 1)):
            if candidate == config.HEADER_ROW:
                continue
            hits = sum(1 for c in rows[candidate - 1]
                       if _normalise(c) in config.FIELD_BY_HEADER)
            if hits >= 3:
                _line(WARN, f"row {candidate} looks more like the header ({hits} known "
                            f"columns) — set config.HEADER_ROW = {candidate} and "
                            f"FIRST_DATA_ROW = {candidate + 1}")


def _inspect_data(rows: list[tuple]) -> None:
    """The data rows: entity codes, and the types that the reader insists on."""
    header = rows[config.HEADER_ROW - 1] if len(rows) >= config.HEADER_ROW else ()
    columns = {}
    for index, cell in enumerate(header):
        field = config.FIELD_BY_HEADER.get(_normalise(cell))
        if field is not None and field not in columns:
            columns[field] = index

    body = rows[config.FIRST_DATA_ROW - 1:]
    if not body:
        _line(BAD, f"no rows at or after row {config.FIRST_DATA_ROW}")
        return

    def value(row, field):
        index = columns.get(field)
        if index is None or index >= len(row):
            return None
        return row[index]

    # Entity codes. This is the trap that produces an EMPTY table rather than an error:
    # a row counts as a detail row only if its company code is a declared entity.
    if "company_code" in columns:
        seen = Counter(str(value(r, "company_code")).strip()
                       for r in body if isinstance(value(r, "company_code"), str)
                       and str(value(r, "company_code")).strip())
        known = {c: n for c, n in seen.items() if c in config.ENTITY_LABEL}
        unknown = {c: n for c, n in seen.items() if c not in config.ENTITY_LABEL}
        if known:
            _line(OK, f"entity codes recognised: "
                      f"{dict(sorted(known.items(), key=lambda kv: -kv[1]))}")
        else:
            _line(BAD, "NO recognised entity codes in the first 200 rows")
        if unknown:
            status = BAD if not known else WARN
            _line(status, f"unrecognised codes in the Company Code column: "
                          f"{dict(sorted(unknown.items(), key=lambda kv: -kv[1])[:8])}")
            _note("-> every row carrying these is SKIPPED silently. If they are real")
            _note("   entities, add them to config.ENTITY_LABEL and ENTITY_GROUP.")

    detail = [r for r in body
              if isinstance(value(r, "company_code"), str)
              and str(value(r, "company_code")).strip() in config.ENTITY_LABEL]
    _line(OK if detail else BAD,
          f"{len(detail)} of {len(body)} scanned rows would be read as detail rows")
    if not detail:
        return

    sample = detail[0]

    # Customer must be TEXT. Read as a number it loses leading zeros, and every join
    # against those customers then misses them silently — so ingest raises instead.
    customer = value(sample, "customer")
    if isinstance(customer, str):
        _line(OK, f"Customer is text, e.g. {customer!r}")
    else:
        _line(BAD, f"Customer is {type(customer).__name__} ({customer!r}), not text")
        _note("-> ingest RAISES on this, deliberately: as a number the value has")
        _note("   already lost any leading zero. Fix the export to write the column")
        _note("   as text, or format it as Text in Excel and re-save.")

    when = value(sample, "document_date")
    if isinstance(when, (datetime, date)):
        _line(OK, f"Document Date is a real date, e.g. {when}")
    else:
        _line(BAD, f"Document Date is {type(when).__name__} ({when!r}), not a date")
        _note("-> the column is stored as text. Format it as a date and re-save.")

    for field in ("amount_original", "amount_local"):
        raw = value(sample, field)
        if isinstance(raw, bool) or not isinstance(raw, (int, float)):
            _line(BAD, f"{field} is {type(raw).__name__} ({raw!r}), not a number")
        else:
            _line(OK, f"{field} is numeric, e.g. {raw:,.2f}")

    currencies = Counter(str(value(r, "local_currency")) for r in detail)
    _line(OK if set(currencies) <= {"TRY"} else WARN,
          f"local currency values: {dict(currencies)}")
    if set(currencies) - {"TRY"}:
        _note("-> the pipeline sums amount_local as one number. Mixed local")
        _note("   currencies would be summing different units.")

    # Unapplied cash: the invoice column holds a Turkish status note instead of a number.
    # Invoice references are a single unbroken token, so a value containing a space is
    # prose rather than a reference — that, not "starts with a letter", is the tell.
    notes = Counter(
        text for r in detail
        if isinstance(value(r, "government_invoice"), str)
        and (" " in (text := str(value(r, "government_invoice")).strip())
             or text in config.STATUS_NOTES))
    if notes:
        unknown = {n: c for n, c in notes.items() if n not in config.STATUS_NOTES}
        _line(OK if not unknown else WARN,
              f"unapplied-cash notes seen: {dict(list(notes.items())[:5])}")
        if unknown:
            _note(f"-> {sorted(unknown)[:5]} are not in config.STATUS_NOTES. Add them")
            _note("   so they are counted as unapplied cash. Diagnostics only — the")
            _note("   money is included in the totals either way.")
    else:
        _line(OK, "no unapplied-cash notes in the scanned rows")


def main(argv: list[str]) -> int:
    target = argv[1] if len(argv) > 1 else config.CORPUS_DIR
    print("=" * 74)
    print("preflight — does this data match what config.py declares?")
    print("=" * 74)

    if os.path.isfile(target):
        inspect_workbook(target)
    else:
        paths = inspect_folder(target)
        if paths:
            # First, middle and last: layout drift across a long corpus is real, and a
            # column inserted halfway through would pass a check on file one alone.
            picks = {0: "first", len(paths) // 2: "middle", len(paths) - 1: "last"}
            for index, label in sorted(picks.items()):
                print(f"\n--- {label} file")
                inspect_workbook(paths[index])

    print("\n" + "=" * 74)
    print("Read every FAIL, then edit config.py. See HANDOVER.md.")
    print("=" * 74)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
